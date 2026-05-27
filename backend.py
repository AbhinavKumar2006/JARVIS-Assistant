import sys
import time
import subprocess
import requests
from pathlib import Path
import mimetypes

from brain import ask_llama, save_history

OLLAMA_URL = "http://localhost:11434"
MAX_FILE_CHARS = 20000

TEXT_EXTENSIONS = {
    ".txt", ".py", ".java", ".html", ".css", ".js", ".json", ".md", ".csv",
    ".cpp", ".c", ".h", ".hpp", ".xml", ".yml", ".yaml", ".bat", ".ps1",
    ".sql", ".log", ".ini", ".cfg", ".properties"
}
IMAGE_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".tiff", ".webp"}


def is_ollama_running():
    try:
        response = requests.get(OLLAMA_URL, timeout=3)
        return response.status_code == 200
    except Exception:
        return False


def start_ollama():
    try:
        if sys.platform.startswith("win"):
            subprocess.Popen(
                ["ollama", "serve"],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL,
                creationflags=subprocess.CREATE_NO_WINDOW
            )
        else:
            subprocess.Popen(
                ["ollama", "serve"],
                stdout=subprocess.DEVNULL,
                stderr=subprocess.DEVNULL
            )

        return True

    except Exception:
        return False


def wait_for_ollama(max_wait=20):
    for _ in range(max_wait):
        if is_ollama_running():
            return True
        time.sleep(1)
    return False


def cut_text(text, limit=MAX_FILE_CHARS):
    text = text or ""
    if len(text) <= limit:
        return text
    return text[:limit] + "\n\n[File was long, so only the first part was sent to Llama.]"


def read_text_file(path: Path):
    encodings = ["utf-8", "utf-8-sig", "cp1252", "latin-1"]
    last_error = None
    for enc in encodings:
        try:
            return path.read_text(encoding=enc, errors="strict")
        except Exception as e:
            last_error = e
    return path.read_text(encoding="utf-8", errors="ignore") or f"Could not decode text cleanly: {last_error}"


def extract_pdf(path: Path):
    try:
        from pypdf import PdfReader
    except Exception:
        try:
            from PyPDF2 import PdfReader
        except Exception:
            return "PDF support needs pypdf. Install it with: pip install pypdf"

    reader = PdfReader(str(path))
    parts = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            parts.append(f"\n--- Page {i} ---\n{text}")
    if not parts:
        return "No selectable text was found in this PDF. It may be a scanned PDF image. OCR is needed."
    return "".join(parts)


def extract_docx(path: Path):
    try:
        import docx
    except Exception:
        return "DOCX support needs python-docx. Install it with: pip install python-docx"
    doc = docx.Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def extract_pptx(path: Path):
    try:
        from pptx import Presentation
    except Exception:
        return "PPTX support needs python-pptx. Install it with: pip install python-pptx"
    prs = Presentation(str(path))
    parts = []
    for idx, slide in enumerate(prs.slides, start=1):
        lines = []
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text.strip():
                lines.append(shape.text.strip())
        if lines:
            parts.append(f"\n--- Slide {idx} ---\n" + "\n".join(lines))
    return "".join(parts) if parts else "No readable text found in this PowerPoint."


def extract_xlsx(path: Path):
    try:
        import openpyxl
    except Exception:
        return "Excel support needs openpyxl. Install it with: pip install openpyxl"
    wb = openpyxl.load_workbook(str(path), data_only=True, read_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"\n--- Sheet: {ws.title} ---")
        for row in ws.iter_rows(max_row=80, max_col=20, values_only=True):
            values = ["" if v is None else str(v) for v in row]
            if any(v.strip() for v in values):
                parts.append(" | ".join(values))
    return "\n".join(parts)


def extract_image_text(path: Path):
    try:
        from PIL import Image
        import pytesseract
    except Exception:
        return "Image OCR needs pillow and pytesseract. Install them and Tesseract OCR to read text from images."
    try:
        text = pytesseract.image_to_string(Image.open(path))
    except Exception as e:
        return "Could not OCR this image: " + str(e)
    if text.strip():
        return text
    return "No readable text was found in this image. A text-only Llama model cannot visually understand image objects."


def extract_file_content(path: Path):
    ext = path.suffix.lower()
    mime, _ = mimetypes.guess_type(str(path))

    if ext in TEXT_EXTENSIONS or (mime and mime.startswith("text/")):
        content = read_text_file(path)
        kind = "text/code file"
    elif ext == ".pdf":
        content = extract_pdf(path)
        kind = "PDF file"
    elif ext == ".docx":
        content = extract_docx(path)
        kind = "Word document"
    elif ext == ".pptx":
        content = extract_pptx(path)
        kind = "PowerPoint file"
    elif ext in [".xlsx", ".xlsm"]:
        content = extract_xlsx(path)
        kind = "Excel spreadsheet"
    elif ext in IMAGE_EXTENSIONS:
        content = extract_image_text(path)
        kind = "image file"
    else:
        size_kb = path.stat().st_size / 1024
        kind = "unsupported/binary file"
        content = (
            "I could not safely extract readable text from this file type.\n"
            f"File extension: {ext or 'none'}\n"
            f"Detected MIME type: {mime or 'unknown'}\n"
            f"File size: {size_kb:.2f} KB\n"
            "Jarvis can still explain what this file type usually is, but not its internal content."
        )

    return kind, cut_text(content)


def build_file_prompt(raw_request: str):
    # Protocol sent by Java:
    # __JARVIS_FILE_UPLOAD__\nFILE_PATH=C:\...\file.pdf\nUSER_TASK:\nexplain this
    lines = raw_request.splitlines()
    file_path = ""
    task_lines = []
    in_task = False

    for line in lines[1:]:
        if line.startswith("FILE_PATH="):
            file_path = line[len("FILE_PATH="):].strip()
        elif line.strip() == "USER_TASK:":
            in_task = True
        elif in_task:
            task_lines.append(line)

    task = "\n".join(task_lines).strip() or "Explain this file clearly and tell the important points."
    path = Path(file_path)

    if not path.exists():
        return f"The selected file was not found on disk: {file_path}"

    kind, extracted = extract_file_content(path)

    return f"""
The user attached a file in the Jarvis app.

File name: {path.name}
File path: {path}
Detected type: {kind}

User task:
{task}

Extracted/readable file content:
{extracted}

Now perform the user's task using the extracted content. If content could not be extracted, clearly say what is missing and what the user should install or provide.
""".strip()


def main():
    user_question = " ".join(sys.argv[1:]).strip()

    if not user_question:
        user_question = sys.stdin.read().strip()

    if not user_question:
        print("No question received.")
        sys.exit()

    if user_question.startswith("__JARVIS_FILE_UPLOAD__"):
        user_question = build_file_prompt(user_question)

    if not is_ollama_running():
        started = start_ollama()

        if not started:
            print("Failed to start Ollama.")
            sys.exit()

        if not wait_for_ollama():
            print("Ollama started but is not responding.")
            sys.exit()

    ai_reply = ask_llama(user_question)
    save_history(user_question, ai_reply)
    print(ai_reply)


if __name__ == "__main__":
    main()
